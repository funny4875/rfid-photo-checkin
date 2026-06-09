from __future__ import annotations

import csv
import io
import re
import secrets
import shutil
import threading
import time
from collections import defaultdict, deque
from datetime import datetime, timedelta
from pathlib import Path, PurePosixPath
from urllib.parse import quote

from flask import Flask, abort, jsonify, render_template, request, send_file, send_from_directory
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
ROOT_DIR = BASE_DIR.parent
VERSION_FILE = ROOT_DIR / "VERSION"
STUDENT_FILE = BASE_DIR / "student_data.txt"
BACKUP_DIR = BASE_DIR / "backups"
STUDENT_BACKUP_FILE = BACKUP_DIR / "student_data_previous.txt"
STUDENT_RESTORE_LOCK = BACKUP_DIR / "student_data_restore_used.flag"
LOCATION_FILE = BASE_DIR / "場域對應.txt"
PHOTO_DIR_FILE = BASE_DIR / "照片資料夾.txt"
NO_PICTURE = "noPicture.jpg"
VALID_DIRECTIONS = {"刷進", "刷出"}
RECORD_RE = re.compile(r"^門禁記錄(?:\d+)?_(\d{8})\.txt$")
STUDENT_HEADERS = ["編號", "UID", "學號", "座號", "班級", "姓名"]
PHOTO_EXTENSIONS = {".jpg", ".jpeg", ".png"}
PHOTO_TOKEN_TTL_SECONDS = 60
PHOTO_TOKEN_MAX_USES = 3
PHOTO_RATE_WINDOW_SECONDS = 60
PHOTO_RATE_MAX_REQUESTS = 120
PHOTO_RATE_BLOCK_SECONDS = 300

app = Flask(__name__)
app.config["TEMPLATES_AUTO_RELOAD"] = True
_file_lock = threading.Lock()
_photo_lock = threading.Lock()
_photo_tokens: dict[str, dict[str, object]] = {}
_photo_hits: defaultdict[str, deque[float]] = defaultdict(deque)
_photo_blocked_until: dict[str, float] = {}


def today_str() -> str:
    return datetime.now().strftime("%Y%m%d")


def app_version() -> str:
    if VERSION_FILE.exists():
        return VERSION_FILE.read_text(encoding="utf-8").strip()
    return "0.0.0"


def ensure_photo_dir_file() -> None:
    if PHOTO_DIR_FILE.exists():
        return
    PHOTO_DIR_FILE.write_text("ccsh_data\n", encoding="utf-8")


def load_photo_dir() -> Path:
    ensure_photo_dir_file()
    value = PHOTO_DIR_FILE.read_text(encoding="utf-8-sig").strip()
    if not value:
        value = "ccsh_data"
    path = Path(value)
    if not path.is_absolute():
        path = BASE_DIR / path
    return path


def record_path(machine_id: str, date_text: str | None = None) -> Path:
    return BASE_DIR / f"門禁記錄{machine_id}_{date_text or today_str()}.txt"


def merged_record_path(date_text: str | None = None) -> Path:
    return BASE_DIR / f"門禁記錄_{date_text or today_str()}.txt"


def ensure_location_file() -> None:
    if LOCATION_FILE.exists():
        return
    LOCATION_FILE.write_text("0\t警衛室\n1\t教官室\n", encoding="utf-8")


def load_locations() -> list[dict[str, str]]:
    ensure_location_file()
    locations: list[dict[str, str]] = []
    for line in LOCATION_FILE.read_text(encoding="utf-8-sig").splitlines():
        line = line.strip()
        if not line:
            continue
        parts = line.split("\t", 1)
        machine_id = parts[0].strip()
        label = (parts[1] if len(parts) > 1 else f"機台{machine_id}").strip()
        if machine_id.isdigit() and label:
            locations.append({"machine_id": machine_id, "label": label})
    if not locations:
        locations = [{"machine_id": "0", "label": "警衛室"}, {"machine_id": "1", "label": "教官室"}]
        save_locations(locations)
    return sorted(locations, key=lambda item: int(item["machine_id"]))


def save_locations(locations: list[dict[str, str]]) -> None:
    seen: set[str] = set()
    normalized: list[dict[str, str]] = []
    for location in locations:
        machine_id = str(location.get("machine_id") or "").strip()
        label = str(location.get("label") or "").strip()
        if not machine_id.isdigit() or not label or machine_id in seen:
            continue
        seen.add(machine_id)
        normalized.append({"machine_id": machine_id, "label": label})
    if not normalized:
        raise ValueError("至少需要一個場域，且機台編號需為數字")
    normalized.sort(key=lambda item: int(item["machine_id"]))
    text = "".join(f"{item['machine_id']}\t{item['label']}\n" for item in normalized)
    with _file_lock:
        LOCATION_FILE.write_text(text, encoding="utf-8")


def valid_machine_ids() -> set[str]:
    return {location["machine_id"] for location in load_locations()}


def format_seat(value: str) -> str:
    value = (value or "").strip()
    if not value:
        return ""
    return value if value.endswith("號") else f"{value}號"


def display_uid(value: str) -> str:
    value = (value or "").strip()
    return "?" if not value or value == "?????:?????" else value


def photo_filename(student_id: str) -> str:
    photo_dir = load_photo_dir()
    for suffix in (".jpg", ".JPG"):
        candidate = photo_dir / f"{student_id}{suffix}"
        if candidate.exists():
            return f"ccsh_data/{candidate.name}"
    return NO_PICTURE


def cleanup_photo_guards(now: float | None = None) -> None:
    current = now or time.time()
    expired_tokens = [
        token
        for token, token_data in _photo_tokens.items()
        if float(token_data.get("expires_at") or 0) <= current
    ]
    for token in expired_tokens:
        _photo_tokens.pop(token, None)
    expired_blocks = [
        client_id
        for client_id, blocked_until in _photo_blocked_until.items()
        if blocked_until <= current
    ]
    for client_id in expired_blocks:
        _photo_blocked_until.pop(client_id, None)


def photo_client_id() -> str:
    return request.remote_addr or "unknown"


def issue_photo_url(photo_file: str) -> str:
    if not photo_file.startswith("ccsh_data/"):
        return f"/photo/{NO_PICTURE}"
    filename = PurePosixPath(photo_file.replace("\\", "/")).name
    token = secrets.token_urlsafe(24)
    with _photo_lock:
        cleanup_photo_guards()
        _photo_tokens[token] = {
            "filename": filename,
            "expires_at": time.time() + PHOTO_TOKEN_TTL_SECONDS,
            "uses_left": PHOTO_TOKEN_MAX_USES,
        }
    return f"/photo/ccsh_data/{quote(filename)}?token={quote(token)}"


def student_for_response(student: dict[str, str]) -> dict[str, str]:
    public_student = dict(student)
    public_student["photo_url"] = issue_photo_url(student.get("photo_file") or NO_PICTURE)
    public_student.pop("photo_file", None)
    return public_student


def consume_photo_token(token: str | None, filename: str) -> bool:
    if not token:
        return False
    with _photo_lock:
        cleanup_photo_guards()
        token_data = _photo_tokens.get(token)
        if not token_data or token_data.get("filename") != filename:
            return False
        uses_left = int(token_data.get("uses_left") or 0)
        if uses_left <= 0:
            _photo_tokens.pop(token, None)
            return False
        if uses_left == 1:
            _photo_tokens.pop(token, None)
        else:
            token_data["uses_left"] = uses_left - 1
        return True


def allow_photo_request() -> bool:
    now = time.time()
    client_id = photo_client_id()
    with _photo_lock:
        cleanup_photo_guards(now)
        blocked_until = _photo_blocked_until.get(client_id)
        if blocked_until and blocked_until > now:
            return False
        hits = _photo_hits[client_id]
        while hits and hits[0] <= now - PHOTO_RATE_WINDOW_SECONDS:
            hits.popleft()
        hits.append(now)
        if len(hits) > PHOTO_RATE_MAX_REQUESTS:
            _photo_blocked_until[client_id] = now + PHOTO_RATE_BLOCK_SECONDS
            hits.clear()
            return False
    return True


def excel_cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def create_student_template_workbook() -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "學生資料"
    sheet.append(STUDENT_HEADERS)
    sheet.append(["1", "01234:56789", "400001", "1", "高一1", "王小明"])
    sheet.append(["2", "02345:67890", "400002", "2", "高一1", "陳小華"])

    header_fill = PatternFill("solid", fgColor="12736B")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")

    widths = [10, 16, 14, 10, 14, 16]
    for index, width in enumerate(widths, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = width
        for cell in sheet[letter]:
            cell.number_format = "@"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:F3"
    note = workbook.create_sheet("說明")
    note["A1"] = "匯入說明"
    note["A1"].font = Font(bold=True, size=14)
    notes = [
        "請保留「學生資料」工作表第一列欄位名稱。",
        "所有欄位建議使用文字格式，避免學號或 UID 前導 0 被 Excel 移除。",
        "UID 格式範例：01234:56789；若尚未建檔可留空或填 ?????:?????。",
        "匯入後會覆寫 server/student_data.txt，系統會自動保留最近一次舊檔備份。",
    ]
    for row_index, text in enumerate(notes, start=2):
        note.cell(row=row_index, column=1, value=text)
    note.column_dimensions["A"].width = 90
    return workbook


def backup_current_student_file() -> None:
    BACKUP_DIR.mkdir(exist_ok=True)
    for old_backup in BACKUP_DIR.glob("student_data_*.txt"):
        if old_backup != STUDENT_BACKUP_FILE:
            old_backup.unlink(missing_ok=True)
    if STUDENT_FILE.exists():
        shutil.copy2(STUDENT_FILE, STUDENT_BACKUP_FILE)
        STUDENT_RESTORE_LOCK.unlink(missing_ok=True)
    else:
        STUDENT_BACKUP_FILE.unlink(missing_ok=True)
        STUDENT_RESTORE_LOCK.write_text("no-backup\n", encoding="utf-8")


def student_restore_status() -> dict[str, object]:
    backup_exists = STUDENT_BACKUP_FILE.exists()
    already_restored = STUDENT_RESTORE_LOCK.exists()
    return {
        "can_restore": backup_exists and not already_restored,
        "backup_exists": backup_exists,
        "already_restored": already_restored,
    }


def restore_previous_student_file() -> dict[str, object]:
    status = student_restore_status()
    if not status["backup_exists"]:
        raise ValueError("目前沒有上一版學生資料可回復")
    if status["already_restored"]:
        raise ValueError("已回復上一版資料，需重新上傳後才能再次回復")
    with _file_lock:
        shutil.copy2(STUDENT_BACKUP_FILE, STUDENT_FILE)
        STUDENT_RESTORE_LOCK.write_text(datetime.now().isoformat(timespec="seconds"), encoding="utf-8")
    return student_restore_status()


def save_uploaded_student_workbook(file_storage) -> dict[str, object]:
    workbook = load_workbook(file_storage, data_only=True, read_only=True)
    sheet = workbook["學生資料"] if "學生資料" in workbook.sheetnames else workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = [excel_cell_text(value) for value in next(rows)]
    except StopIteration:
        raise ValueError("Excel 檔案沒有資料")

    missing = [header for header in STUDENT_HEADERS if header not in header_row]
    if missing:
        raise ValueError(f"缺少欄位：{', '.join(missing)}")

    indexes = {header: header_row.index(header) for header in STUDENT_HEADERS}
    output_rows: list[dict[str, str]] = []
    for row in rows:
        values = list(row)
        item = {
            header: excel_cell_text(values[indexes[header]]) if indexes[header] < len(values) else ""
            for header in STUDENT_HEADERS
        }
        if not any(item.values()):
            continue
        if not item["學號"]:
            continue
        output_rows.append(item)

    if not output_rows:
        raise ValueError("沒有可匯入的學生資料")

    backup_current_student_file()

    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=STUDENT_HEADERS, delimiter="\t", lineterminator="\n")
    writer.writeheader()
    writer.writerows(output_rows)
    with _file_lock:
        STUDENT_FILE.write_text(buffer.getvalue(), encoding="utf-8")
    return {"count": len(output_rows)}


def load_students() -> dict[str, dict[str, str]]:
    students: dict[str, dict[str, str]] = {}
    if not STUDENT_FILE.exists():
        return students

    with STUDENT_FILE.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        for row in reader:
            student_id = (row.get("學號") or "").strip()
            if not student_id:
                continue
            students[student_id] = {
                "number": (row.get("編號") or "").strip(),
                "uid": (row.get("UID") or "").strip(),
                "record_uid": display_uid(row.get("UID") or ""),
                "student_id": student_id,
                "seat": format_seat(row.get("座號") or ""),
                "class_name": (row.get("班級") or "").strip(),
                "name": (row.get("姓名") or "").strip(),
                "photo_file": photo_filename(student_id),
            }
    return students


def load_students_by_uid() -> dict[str, dict[str, str]]:
    return {
        student["uid"]: student
        for student in load_students().values()
        if student.get("uid") and student.get("uid") != "?????:?????"
    }


def get_student_or_none(student_id: str) -> dict[str, str] | None:
    return load_students().get((student_id or "").strip())


def get_student_by_uid_or_none(uid: str) -> dict[str, str] | None:
    return load_students_by_uid().get((uid or "").strip())


def parse_record(line: str, index: int) -> dict[str, str | int]:
    parts = line.rstrip("\n").split("\t")
    while len(parts) < 7:
        parts.append("")
    return {
        "index": index,
        "raw": line.rstrip("\n"),
        "time": parts[0],
        "uid": parts[1],
        "student_id": parts[2],
        "class_name": parts[3],
        "seat": parts[4],
        "name": parts[5],
        "direction": parts[6],
    }


def read_record_lines(path: Path) -> list[str]:
    if not path.exists():
        return []
    return path.read_text(encoding="utf-8").splitlines()


def records_for_machine(machine_id: str, date_text: str | None = None) -> list[dict[str, str | int]]:
    return [parse_record(line, index) for index, line in enumerate(read_record_lines(record_path(machine_id, date_text)))]


def write_record(machine_id: str, student: dict[str, str], direction: str) -> dict[str, str | int]:
    now = datetime.now().strftime("%H:%M:%S")
    line = "\t".join(
        [
            now,
            student["record_uid"],
            student["student_id"],
            student["class_name"],
            student["seat"],
            student["name"],
            direction,
        ]
    )
    path = record_path(machine_id)
    with _file_lock:
        with path.open("a", encoding="utf-8", newline="") as handle:
            handle.write(line + "\n")
        index = len(read_record_lines(path)) - 1
    return parse_record(line, index)


def merge_today_records(date_text: str | None = None) -> dict[str, str | int]:
    date_value = date_text or today_str()
    machine_paths = sorted(BASE_DIR.glob(f"門禁記錄[0-9]*_{date_value}.txt"))
    lines: list[str] = []
    for path in machine_paths:
        lines.extend(read_record_lines(path))

    lines.sort(key=lambda line: line.split("\t", 1)[0] if line else "")
    output_path = merged_record_path(date_value)
    with _file_lock:
        output_path.write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8")
    return {"file": output_path.name, "count": len(lines)}


def archive_old_records(days: int = 7) -> dict[str, object]:
    cutoff = datetime.now().date() - timedelta(days=days)
    moved: list[str] = []
    with _file_lock:
        for path in BASE_DIR.glob("門禁記錄*.txt"):
            match = RECORD_RE.match(path.name)
            if not match:
                continue
            record_date = datetime.strptime(match.group(1), "%Y%m%d").date()
            if record_date > cutoff:
                continue
            year_dir = BASE_DIR / str(record_date.year)
            year_dir.mkdir(exist_ok=True)
            destination = year_dir / path.name
            shutil.move(str(path), str(destination))
            moved.append(path.name)
    return {"moved": moved, "count": len(moved)}


def delete_today_record(machine_id: str, index: int) -> dict[str, object]:
    path = record_path(machine_id)
    with _file_lock:
        lines = read_record_lines(path)
        if index < 0 or index >= len(lines):
            abort(404, description="找不到選取的今日記錄")
        removed = parse_record(lines.pop(index), index)
        path.write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8")
    return {"deleted": removed, "records": records_for_machine(machine_id)}


def scheduler_loop() -> None:
    last_merge_date = ""
    last_archive_date = ""
    while True:
        now = datetime.now()
        current_date = now.strftime("%Y%m%d")
        if now.strftime("%H:%M") == "12:10" and last_merge_date != current_date:
            merge_today_records(current_date)
            last_merge_date = current_date
        if now.strftime("%H:%M") == "00:15" and last_archive_date != current_date:
            archive_old_records()
            last_archive_date = current_date
        time.sleep(30)


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/admin")
def admin():
    return render_template("admin.html")


@app.route("/photo/<path:filename>")
def photo(filename: str):
    if filename == "noUID.jpg":
        return send_from_directory(BASE_DIR, "noUID.jpg")
    if filename == NO_PICTURE:
        return send_from_directory(BASE_DIR, NO_PICTURE)

    photo_path = PurePosixPath(filename.replace("\\", "/"))
    if (
        len(photo_path.parts) == 2
        and photo_path.parts[0] == "ccsh_data"
        and ".." not in photo_path.parts
        and photo_path.suffix.lower() in PHOTO_EXTENSIONS
        and not photo_path.name.lower().startswith("student_data")
    ):
        if not allow_photo_request():
            abort(429)
        if not consume_photo_token(request.args.get("token"), photo_path.name):
            abort(404)
        return send_from_directory(load_photo_dir(), photo_path.name, max_age=0)
    abort(404)


@app.route("/favicon.ico")
def favicon():
    return send_from_directory(BASE_DIR / "static", "favicon.ico")


@app.get("/api/student/<student_id>")
def api_student(student_id: str):
    student = get_student_or_none(student_id)
    if not student:
        return jsonify({"error": "查無此學號"}), 404
    return jsonify({"student": student_for_response(student)})


@app.post("/api/checkin")
def api_checkin():
    payload = request.get_json(silent=True) or request.form
    student_id = (payload.get("student_id") or "").strip()
    machine_id = str(payload.get("machine_id") or "0")
    direction = payload.get("direction") or "刷進"
    if machine_id not in valid_machine_ids() or direction not in VALID_DIRECTIONS:
        return jsonify({"error": "機台或刷進刷出選項錯誤"}), 400
    student = get_student_or_none(student_id)
    if not student:
        return jsonify({"error": "查無此學號"}), 404
    record = write_record(machine_id, student, direction)
    return jsonify({"student": student_for_response(student), "record": record, "records": records_for_machine(machine_id)})


@app.post("/api/checkin_uid")
def api_checkin_uid():
    payload = request.get_json(silent=True) or request.form
    uid = (payload.get("uid") or "").strip()
    machine_id = str(payload.get("machine_id") or "0")
    direction = payload.get("direction") or "刷進"
    if machine_id not in valid_machine_ids() or direction not in VALID_DIRECTIONS:
        return jsonify({"error": "機台或刷進刷出選項錯誤"}), 400
    if not uid:
        return jsonify({"error": "沒有收到 UID"}), 400
    student = get_student_by_uid_or_none(uid)
    if not student:
        return jsonify({"error": "UID 未建檔", "uid": uid, "records": records_for_machine(machine_id)}), 404
    record = write_record(machine_id, student, direction)
    return jsonify({"student": student_for_response(student), "record": record, "records": records_for_machine(machine_id), "uid": uid})


@app.get("/api/rfid/status")
def api_rfid_status():
    return jsonify({"enabled": False, "error": "RFID 讀卡程式請由前端主機 client.bat 啟動"})


@app.get("/api/locations")
def api_locations():
    return jsonify({"locations": load_locations()})


@app.get("/api/version")
def api_version():
    return jsonify({"version": app_version()})


@app.post("/api/admin/locations")
def api_admin_locations():
    payload = request.get_json(silent=True) or {}
    locations = payload.get("locations") or []
    if not isinstance(locations, list):
        return jsonify({"error": "場域資料格式錯誤"}), 400
    try:
        save_locations(locations)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    return jsonify({"locations": load_locations()})


@app.get("/api/records")
def api_records():
    machine_id = request.args.get("machine_id", "0")
    if machine_id not in valid_machine_ids():
        return jsonify({"error": "機台錯誤"}), 400
    return jsonify({"records": records_for_machine(machine_id)})


@app.delete("/api/records/<machine_id>/<int:index>")
def api_delete_record(machine_id: str, index: int):
    if machine_id not in valid_machine_ids():
        return jsonify({"error": "機台錯誤"}), 400
    return jsonify(delete_today_record(machine_id, index))


@app.get("/api/admin/records")
def api_admin_records():
    locations = load_locations()
    return jsonify(
        {
            "locations": locations,
            "machines": {location["machine_id"]: records_for_machine(location["machine_id"]) for location in locations},
            "merged": [parse_record(line, index) for index, line in enumerate(read_record_lines(merged_record_path()))],
        }
    )


@app.post("/api/admin/merge")
def api_admin_merge():
    return jsonify(merge_today_records())


@app.post("/api/admin/archive")
def api_admin_archive():
    return jsonify(archive_old_records())


@app.get("/api/admin/student-template")
def api_admin_student_template():
    workbook = create_student_template_workbook()
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name="student_data_template.xlsx",
    )


@app.post("/api/admin/student-data/upload")
def api_admin_student_data_upload():
    upload = request.files.get("student_file")
    if not upload or not upload.filename:
        return jsonify({"error": "請選擇 Excel 檔案"}), 400
    if not upload.filename.lower().endswith(".xlsx"):
        return jsonify({"error": "只支援 .xlsx 檔案"}), 400
    try:
        result = save_uploaded_student_workbook(upload)
    except ValueError as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception as exc:
        return jsonify({"error": f"匯入失敗：{exc}"}), 400
    return jsonify({"message": "學生資料已更新", **result})


@app.get("/api/admin/student-data/restore/status")
def api_admin_student_data_restore_status():
    return jsonify(student_restore_status())


@app.post("/api/admin/student-data/restore")
def api_admin_student_data_restore():
    try:
        status = restore_previous_student_file()
    except ValueError as exc:
        return jsonify({"error": str(exc), **student_restore_status()}), 400
    except Exception as exc:
        return jsonify({"error": f"回復失敗：{exc}", **student_restore_status()}), 400
    return jsonify({"message": "已回復上一版學生資料", **status})


if __name__ == "__main__":
    threading.Thread(target=scheduler_loop, daemon=True).start()
    app.run(host="0.0.0.0", port=5000, debug=False)
